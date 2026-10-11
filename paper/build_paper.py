"""Build the NavCert Offline IEEE conference paper (DOCX) and convert it to PDF.

Run:  ml\\.venv\\Scripts\\python.exe paper\\build_paper.py
Figures come from paper/make_figures.py (run it first). Every number in the text is taken from
ml/ivp_eval/README.md, results.json, block_matching.json, paper/figures/derived_stats.json,
ml/README.md, ml/ai_eval/README.md or the app source code.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from omml import bar, cat, delim, frac, func, lim_low, math, nary, sqrt, sub, subsup, sup, t, v  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FIG = HERE / "figures"
DOCX = HERE / "NavCert_Offline_IEEE.docx"
PDF = HERE / "NavCert_Offline_IEEE.pdf"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
TITLE = "NavCert Offline: Classical Image Processing for Robust On-Device Digitization of Maritime Certificates"

COL_W = 3.5      # inches
FULL_W = 7.16
BODY_PT = 10
STATS = json.loads((FIG / "derived_stats.json").read_text())


def pct(x, d=1):
    return f"{x * 100:.{d}f}"


# =====================================================================================
# References: parsed from docs/research/references.md, renumbered by first citation.
# =====================================================================================
def load_references():
    text = (ROOT / "docs" / "research" / "references.md").read_text(encoding="utf-8")
    refs = {}
    for m in re.finditer(r"^\[(\d+)\] (.+?)\s*$", text, re.M):
        body = m.group(2).strip()
        url = None
        mu = re.search(r"\s(https?://.+)$", body)
        if mu:
            url = mu.group(1).strip()
            body = body[: mu.start()].rstrip()
        if not body.endswith("."):
            body += "."
        if "ISBN" in body or "Graphics Gems" in body:
            url = None                             # books: the bibliographic entry is the locator
        if url and url.startswith("https://doi.org/"):
            body += f" doi: {url[len('https://doi.org/'):]}."
        elif url and "arxiv.org" in url:
            pass                                   # the arXiv identifier is already in the entry
        elif url:
            body += f" [Online]. Available: {url}"
        refs[int(m.group(1))] = body
    return refs


# =====================================================================================
# Low-level formatting helpers
# =====================================================================================
def set_font(run, size=BODY_PT, bold=False, italic=False, small_caps=False, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.small_caps = small_caps


def keep_next(p):
    p.paragraph_format.keep_with_next = True


def set_cols(section, n, space_in=0.25):
    sp = section._sectPr
    cols = sp.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sp.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_in * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def page_setup(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.4)


TOKEN = re.compile(r"(\[@[0-9@,\s]+\]|\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\}|\{(?:fig|tab|eq):[a-z0-9_]+\})")


class Paper:
    def __init__(self):
        self.doc = Document()
        self.refs = load_references()
        self.cite_order: list[int] = []
        self.labels: dict[str, int] = {}
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(BODY_PT)
        st.element.get_or_add_rPr()
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        page_setup(self.doc.sections[0])

    # ---------------- citations / labels
    def cite_text(self, inner):
        nums = [int(x) for x in re.findall(r"@(\d+)", inner)]
        new = []
        for n in nums:
            if n not in self.refs:
                raise KeyError(f"reference {n} missing")
            if n not in self.cite_order:
                self.cite_order.append(n)
            new.append(self.cite_order.index(n) + 1)
        new = sorted(set(new))
        parts, i = [], 0
        while i < len(new):
            j = i
            while j + 1 < len(new) and new[j + 1] == new[j] + 1:
                j += 1
            if j - i >= 2:
                parts.append(f"[{new[i]}]–[{new[j]}]")
            else:
                parts.extend(f"[{k}]" for k in new[i:j + 1])
            i = j + 1
        return ", ".join(parts)

    def rich(self, p, text, size=BODY_PT, italic=False, bold=False):
        for tok in TOKEN.split(text):
            if not tok:
                continue
            sub_ = sup_ = False
            it = italic
            if tok.startswith("[@"):
                tok = self.cite_text(tok)
            elif tok.startswith("{") and ":" in tok:
                tok = str(self.labels.get(tok[1:-1], "?"))
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, it = tok[1:-1], not italic
            elif tok.startswith("_{"):
                tok, sub_ = tok[2:-1], True
            elif tok.startswith("^{"):
                tok, sup_ = tok[2:-1], True
            r = p.add_run(tok)
            set_font(r, size, bold=bold, italic=it)
            if sub_:
                r.font.subscript = True
            if sup_:
                r.font.superscript = True
        return p

    # ---------------- blocks
    def para(self, text, indent=True, size=BODY_PT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0):
        p = self.doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.first_line_indent = Inches(0.14) if indent else None
        pf.space_after = Pt(space_after)
        pf.widow_control = True
        return self.rich(p, text, size)

    def heading(self, roman, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        keep_next(p)
        r = p.add_run(f"{roman}. " if roman else "")
        set_font(r, BODY_PT, small_caps=True)
        r = p.add_run(text)
        set_font(r, BODY_PT, small_caps=True)
        return p

    def subheading(self, letter, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        keep_next(p)
        r = p.add_run(f"{letter}. {text}")
        set_font(r, BODY_PT, italic=True)
        return p

    def bullets(self, items, numbered=False):
        for k, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent = Inches(0.22)
            pf.first_line_indent = Inches(-0.14)
            pf.space_after = Pt(1)
            mark = f"{k})" if numbered else "•"
            r = p.add_run(mark + "\u00a0")
            set_font(r, BODY_PT)
            self.rich(p, it)

    def equation(self, label, xml):
        self.eq_n += 1
        self.labels[f"eq:{label}"] = self.eq_n
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        pf.tab_stops.add_tab_stop(Inches(COL_W / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Inches(COL_W - 0.02), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t")
        set_font(r)
        p._p.append(parse_xml(xml))
        r = p.add_run(f"\t({self.eq_n})")
        set_font(r)

    def figure(self, label, path, width, caption):
        self.fig_n += 1
        self.labels[f"fig:{label}"] = self.fig_n
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        keep_next(p)
        p.add_run().add_picture(str(path), width=Inches(width))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.space_before = Pt(2)
        c.paragraph_format.space_after = Pt(6)
        r = c.add_run(f"Fig. {self.fig_n}.\u2002")
        set_font(r, 8)
        self.rich(c, caption, size=8)

    def table(self, label, title, headers, rows, widths, notes=None, size=8, align=None):
        self.tab_n += 1
        roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"][self.tab_n - 1]
        self.labels[f"tab:{label}"] = roman
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        keep_next(p)
        r = p.add_run(f"TABLE {roman}")
        set_font(r, 8)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        keep_next(p)
        r = p.add_run(title)
        set_font(r, 8, small_caps=True)
        tb = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        tbl_pr = tb._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"), ("right", "nil"),
                          ("insideH", "nil"), ("insideV", "nil")):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), val)
            if val == "single":
                el.set(qn("w:sz"), "6"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
            borders.append(el)
        tbl_pr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for edge, w_ in (("top", 10), ("bottom", 10), ("left", 45), ("right", 45)):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), str(w_)); el.set(qn("w:type"), "dxa")
            mar.append(el)
        tbl_pr.append(mar)
        grid = tb._tbl.tblGrid
        for gc, w_ in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w_ * 1440)))
        align = align or (["l"] + ["c"] * (len(headers) - 1))
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
        for ri, row in enumerate([headers] + rows):
            tr = tb.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trpr.append(cs)
            if ri == 0:
                th = OxmlElement("w:tblHeader"); trpr.append(th)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(widths[ci])
                cp = cell.paragraphs[0]
                cp.alignment = amap[align[ci]] if ri else WD_ALIGN_PARAGRAPH.CENTER
                if ri == 0 and ci == 0:
                    cp.alignment = amap[align[0]]
                keep_next(cp)
                self.rich(cp, str(val), size=size)
                if ri == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    tb_ = OxmlElement("w:tcBorders")
                    bt = OxmlElement("w:bottom")
                    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "4"); bt.set(qn("w:color"), "000000")
                    tb_.append(bt); tcpr.append(tb_)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4 if not notes else 0)
        if notes:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            self.rich(p, notes, size=7.5)
        else:
            r = p.add_run(""); set_font(r, 4)

    def section(self, cols):
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        page_setup(sec)
        set_cols(sec, cols)
        return sec


# =====================================================================================
# Content
# =====================================================================================
def build(pre_labels=None):
    P = Paper()
    if pre_labels:
        P.labels.update(pre_labels)
    d = P.doc
    set_cols(d.sections[0], 1)
    S = STATS
    thr = json.loads((ROOT / "ml" / "ivp_eval" / "results.json").read_text())["threshold"]["methods"]
    bmj = json.loads((ROOT / "ml" / "ivp_eval" / "block_matching.json").read_text())

    # ---------------- title block
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(TITLE); set_font(r, 24)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Aadidev Raizada"); set_font(r, 11)
    for line, it in (("Mukesh Patel School of Technology Management and Engineering", True),
                     ("SVKM's NMIMS, Mumbai, India", True), ("aadidevraizada26@gmail.com", False)):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); set_font(r, 10, italic=it)
    p.paragraph_format.space_after = Pt(12)
    P.section(2)

    # ---------------- abstract
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract"); set_font(r, 9, bold=True, italic=True)
    r = p.add_run("—"); set_font(r, 9, bold=True)
    P.rich(p, (
        "Seafarers' certificates are checked in ports and on board, often without a network connection. NavCert "
        "Offline is an Android application that digitizes them entirely on the phone. A classical OpenCV front end "
        "finds the page with Canny edges, morphological closing and polygon approximation, rectifies it with a "
        "four-point homography, screens capture quality with brightness and Laplacian-variance tests, and enhances "
        "it with CLAHE, median filtering and unsharp masking or adaptive thresholding. On-device OCR and a rule, "
        "MiniLM and MobileBERT cascade then extract five certificate fields. We measure the image processing on 200 "
        "paired synthetic camera captures under eight degradations with known ground truth. Rectification lowers the "
        "mean character error rate (CER) from 9.6% to 4.2% and raises five-field exact match from 69.1% to 88.4%; "
        "the paired CER reduction is 5.4 points (95% CI 3.9–7.1). Contrast enhancement gives no average gain with "
        "the evaluation OCR engine, helps under low light and blur, and hurts under overexposure. Under cast "
        "shadows, adaptive thresholding keeps CER below 0.5% where global and Otsu thresholds reach 34–40%. Text "
        "survives 8×8 DCT coding at 0.29 bit/pixel with 0.35% CER, while three-step and 2-D logarithmic block "
        "matching find the true motion vector on only 30–33% of textured text blocks. An in-app Pipeline Lab shows "
        "every operation on the user's own scan."), size=9, bold=True)
    p.paragraph_format.space_after = Pt(4)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Index Terms"); set_font(r, 9, bold=True, italic=True)
    r = p.add_run("—document image processing, perspective rectification, CLAHE, adaptive thresholding, "
                  "mobile OCR, block matching, maritime certificates"); set_font(r, 9, bold=True)
    p.paragraph_format.space_after = Pt(2)

    # ---------------- I. Introduction
    P.heading("I", "Introduction")
    P.para("Merchant shipping runs on certified people. The STCW Convention sets the certification every seafarer "
           "must prove, and its 2010 Manila amendments target fraudulent certificates of competency [@58]. The 2026 "
           "BIMCO/ICS report estimates 2,565,580 STCW-certified seafarers serving 85,148 ships [@62]; India supplies "
           "311,936 of them, 12.16% of the total [@67]. Each certificate is read again and again by employers, "
           "manning agents and port state control officers, and the checks matter: an IMO study around 2000 recorded "
           "12,635 detected forgery cases, with 82% of survey respondents reporting forged certificates of "
           "competency [@60], and Asia-Pacific port state control recorded 877 crew-certificate deficiencies in 2022 "
           "and 1,304 in 2024 [@64]. Electronic certificates help, but IMO guidance assumes that the verifying party "
           "has an Internet connection [@61], which a ship at sea or a remote berth cannot promise.")
    P.para("A hand-held phone photo is a poor input for OCR: the page is a projected trapezium on a cluttered "
           "table, unevenly lit, blurred, noisy and JPEG-compressed. This paper asks how much of that gap classical image and video processing can close "
           "on the device, and measures the answer. Its contributions are:")
    P.bullets([
        "a fully offline Android pipeline that detects, rectifies, checks and enhances a certificate photo, reads it "
        "and extracts five fields (certificate name, issue date, port of issue, issuing authority, expiry date);",
        "a paired, seeded benchmark of 200 synthetic captures that isolates what rectification, enhancement and "
        "binarization contribute, with studies of quality checks, shadow thresholding, DCT coding and block matching;",
        "Pipeline Lab, an in-app tool that runs eleven syllabus-level operations, from point processing to block "
        "matching, on the user's own scan.",
    ], numbered=True)

    # ---------------- II. Related work
    P.heading("II", "Related Work")
    P.para("*Capture and rectification.* OCRdroid guided phone capture but recognized text on a server [@2]; the "
           "ICDAR 2015 SmartDoc competition benchmarked page detection and OCR on smartphone photos [@23]. Classical "
           "rectification fits a homography to four page corners [@48], found with Canny edges [@40], border "
           "following [@46] and Douglas–Peucker simplification [@47]. Learned dewarpers such as DocUNet [@21] and "
           "DocTr [@22] also flatten curled pages but need training data and far more compute.")
    P.para("*Enhancement and binarization.* Histogram equalization is the standard global contrast tool [@37]; "
           "CLAHE equalizes tiles and clips their histograms to limit noise gain [@45]. "
           "Otsu's method picks one global threshold [@39]; Sauvola's method [@50] adapts it locally for unevenly "
           "lit documents. Laplacian variance is a common focus measure [@49].")
    P.para("*Mobile OCR and document models.* Modern line recognizers decode with CTC [@10]. The PP-OCR family "
           "targets small devices [@3], down to a 5M-parameter "
           "PP-OCRv5 [@6] and PP-OCRv6 tiers of 1.5M–34.5M parameters [@8], with a detector based on differentiable "
           "binarization [@11]. ML Kit bundles an offline Latin recognizer [@35]. Layout-aware models [@15,@16] "
           "and OCR-free models such as Donut [@14] are stronger on forms but larger. For extraction, BERT-style "
           "question answering [@27] has compact variants in MobileBERT [@28] and MiniLM [@29]; integer quantization "
           "[@32] and ONNX Runtime Mobile [@33] make them practical on phones. Unlike prior capture systems, this "
           "work keeps every stage on the device and reports a paired measurement of each classical stage.")

    # ---------------- III. System overview
    P.heading("III", "System Overview")
    P.para("Fig. {fig:arch} shows the implemented system. The user photographs a certificate through the system "
           "camera or imports one from the gallery; a crop editor allows manual correction. *DocumentScanner* "
           "prepares the image (Section IV), and ML Kit Text Recognition v2 returns blocks, lines and boxes. In "
           "printed and automatic modes the text is laid out, corrected and passed to the field extractor; in "
           "handwriting mode PP-OCRv5 re-reads each line (Section VI). Documents, JPEG previews (quality 88), folders "
           "and notes stay in app-private files, and the manifest removes the INTERNET and ACCESS_NETWORK_STATE "
           "permissions that the ML Kit dependency would add. The app is written in Kotlin with Jetpack Compose, "
           "targets API 36 (minimum 24) and uses OpenCV 4.13.0 [@36], ML Kit text recognition 16.0.1, ONNX Runtime "
           "Android 1.24.1 and TensorFlow Lite Task Text 0.4.4. Fig. {fig:strip} previews the image processing "
           "stages detailed in Section IV.")
    P.section(1)
    P.figure("arch", FIG / "fig_architecture.png", FULL_W,
             "Implemented architecture. The OpenCV stage serves both recognition paths; correction and field "
             "extraction run only in printed and automatic modes; Pipeline Lab reuses the scanner's stage functions.")
    P.figure("strip", FIG / "fig_pipeline_strip.png", FULL_W,
             "Stages on a synthetic tilted capture: (a) camera frame; (b) closed Canny edges; (c) contours (gray) and "
             "the selected convex quadrilateral; (d) rectification; (e) Auto enhancement; (f) B&W adaptive "
             "threshold. The double border inside the page is a deliberate distractor quadrilateral.")
    P.section(2)
    # ---------------- IV. Image processing pipeline
    P.heading("IV", "Image Processing Pipeline")
    P.para("Table {tab:params} lists every stage with the constants in *DocumentScanner.kt*.")
    P.subheading("A", "Resolution and Page Detection")
    P.para("The app only downsamples: one copy for OCR is capped at 3200 px on its longer side and one for analysis "
           "at 1600 px; smaller photos keep their size. All stages work on luminance [@37],")
    P.equation("gray", math(v("Y"), t("="), t("0.299"), v("R"), t("+"), t("0.587"), v("G"), t("+"), t("0.114"), v("B"), t(".")))
    P.table("params", "Image processing stages and constants",
            ["Stage", "Operation", "Constants"],
            [["Resize", "Bilinear, downscale only", "OCR ≤ 3200, analysis ≤ 1600 px"],
             ["Quality", "Mean gray; var(∇²*f*)", "< 55 dark; > 220 bright; < 65 soft"],
             ["Smooth", "Gaussian blur", "5×5, σ ≈ 1.1"],
             ["Edges", "Canny", "*T*_{L} = 55, *T*_{H} = 170"],
             ["Close", "Morphological closing", "5×5 rectangle"],
             ["Polygon", "Douglas–Peucker", "20 largest; ε = 2% of perimeter"],
             ["Page", "Largest convex 4-gon", "≥ 15% of frame"],
             ["Warp", "Homography, bicubic", "Replicated border"],
             ["Contrast", "CLAHE", "Clip 2.2; 8×8 tiles"],
             ["Denoise", "Median", "3×3"],
             ["Sharpen", "Unsharp mask", "σ/*a* 1.05/1.32 Auto; 1.35/1.50 Clean"],
             ["Binarize", "Adaptive (B&W mode)", "Block 35, *C* = 11"],
             ["Pad", "Constant border", "18 px white"]],
            [0.6, 1.3, 1.6])

    P.para("The analysis image is smoothed with a 5×5 Gaussian kernel,", indent=False)
    P.equation("gauss", math(sub(v("G"), v("σ")), delim(cat(v("x"), t(","), v("y"))), t("="),
                             frac(t("1"), cat(t("2π"), sup(v("σ"), t("2")))),
                             sup(v("e"), cat(t("−"), frac(cat(sup(v("x"), t("2")), t("+"), sup(v("y"), t("2"))),
                                                          cat(t("2"), sup(v("σ"), t("2")))))), t(",")))
    P.para("with σ = 0.3((*k* − 1)/2 − 1) + 0.8 ≈ 1.1 for *k* = 5. Canny's detector [@40] computes the Sobel "
           "gradient", indent=False)
    P.equation("grad", math(v("M"), t("="), sqrt(cat(sup(sub(v("G"), v("x")), t("2")), t("+"), sup(sub(v("G"), v("y")), t("2")))),
                            t(",  "), v("θ"), t("="), func("atan2", delim(cat(sub(v("G"), v("y")), t(","), sub(v("G"), v("x"))))), t(",")))
    P.para("keeps only local maxima of *M* along θ, and links edges by hysteresis: a pixel with *M* > *T*_{H} = 170 "
           "starts an edge that may continue through pixels with *M* > *T*_{L} = 55, so strong page borders seed "
           "edges and the low threshold carries them through shadow. A 5×5 closing [@54],", indent=False)
    P.equation("close", math(v("A"), t("•"), v("B"), t("="), delim(cat(v("A"), t("⊕"), v("B"))), t("⊖"), v("B"), t(",")))
    P.para("bridges breaks in the edge map without growing the outline. Contours are traced [@46], the 20 largest are "
           "simplified by Douglas–Peucker [@47] at 2% of their perimeter, and the largest convex quadrilateral covering "
           "at least 15% of the frame is the page. Its corners are ordered by the extremes of *x* + *y* and *y* − *x* "
           "and scaled to the OCR copy; if no candidate qualifies, the whole frame is used.", indent=False)

    P.subheading("B", "Perspective Rectification")
    P.para("A plane and its image are related by a homography with eight degrees of freedom [@48]:")
    P.equation("homog", math(v("u"), t("="), frac(cat(sub(v("h"), t("11")), v("x"), t("+"), sub(v("h"), t("12")), v("y"), t("+"), sub(v("h"), t("13"))),
                                                   cat(sub(v("h"), t("31")), v("x"), t("+"), sub(v("h"), t("32")), v("y"), t("+"), t("1"))),
                             t(",  "), v("v"), t("="), frac(cat(sub(v("h"), t("21")), v("x"), t("+"), sub(v("h"), t("22")), v("y"), t("+"), sub(v("h"), t("23"))),
                                                         cat(sub(v("h"), t("31")), v("x"), t("+"), sub(v("h"), t("32")), v("y"), t("+"), t("1"))), t(".")))
    P.para("Each corner gives two linear equations, so four corners with no three collinear fix *H* (the direct "
           "linear transform). The target is as wide as the longer of the top and bottom edges and as tall as the "
           "longer side edge, so the aspect ratio is approximate; resampling is bicubic with replicated borders.",
           indent=False)

    P.subheading("C", "Capture Quality Checks")
    P.para("Two features are computed on the whole 1600 px analysis frame before rectification, so the background "
           "counts toward both: the mean gray level and the Laplacian variance [@49]")
    P.equation("lap", math(v("S"), t("="), frac(t("1"), v("N")), nary("∑", cat(v("x"), t(","), v("y")), None,
                                                                          sup(delim(cat(t("∇²"), v("f"), delim(cat(v("x"), t(","), v("y"))), t("−"), v("μ"))), t("2"))),
                           t(",")))
    P.para("with the 4-neighbor Laplacian ∇²*f* and its mean μ. The checks run in order and the first failure is "
           "shown: mean < 55 \"Too dark\", mean > 220 \"Too bright\", *S* < 65 \"Soft image\". The thresholds are "
           "hand-set heuristics.", indent=False)

    P.subheading("D", "Enhancement")
    P.para("Every mode first converts the page to gray. CLAHE [@45] equalizes each of 8×8 tiles through its "
           "cumulative histogram,")
    P.equation("he", math(sub(v("T"), v("k")), delim(v("r")), t("="), delim(cat(v("L"), t("−"), t("1"))),
                          nary("∑", cat(v("j"), t("="), t("0")), v("r"), cat(sub(v("p"), v("k")), delim(v("j")))), t(",")))
    P.para("after clipping the tile histogram at 2.2 times its mean bin height and redistributing the excess, which "
           "caps the slope of *T*_{k} and so the noise gain. Each pixel blends the four nearest tile mappings,",
           indent=False)
    P.equation("blend", math(v("g"), delim(v("r")), t("="), nary("∑", cat(v("i"), t(","), v("j")), None,
                             cat(sub(v("w"), cat(v("i"), v("j"))), sub(v("T"), cat(v("i"), v("j"))), delim(v("r")))), t(",")))
    P.para("with *i*, *j* ∈ {0, 1}, bilinear weights *w*_{00} = (1 − *a*)(1 − *b*), *w*_{10} = *a*(1 − *b*), *w*_{01} = (1 − *a*)*b*, "
           "*w*_{11} = *ab* from the pixel's offsets *a*, *b* ∈ [0, 1] between tile centers, so no seams appear. "
           "Fig. {fig:hist} shows the effect under a shadow: global equalization darkens the shadow further, while "
           "CLAHE keeps both halves legible. A 3×3 median, *g* = median{*f*(*s*,*t*) : (*s*,*t*) ∈ *S*_{xy}}, then "
           "removes speckle without blurring strokes. Auto and Clean modes sharpen by unsharp masking,", indent=False)
    P.figure("hist", FIG / "fig_histogram_clahe.png", COL_W,
             "A shadowed page: rectified gray, global histogram equalization and CLAHE (clip 2.2, 8×8 tiles), with "
             "gray-level histograms.")
    P.equation("usm", math(v("g"), t("="), v("a"), v("f"), t("+"), delim(cat(t("1"), t("−"), v("a"))), sub(v("f"), v("σ")),
                           t("="), v("f"), t("+"), v("k"), delim(cat(v("f"), t("−"), sub(v("f"), v("σ")))), t(",")))
    P.para("with *f*_{σ} a Gaussian blur and *k* = *a* − 1: σ = 1.05, *k* = 0.32 (Auto) or σ = 1.35, *k* = 0.50 "
           "(Clean). B&W mode instead binarizes against a local threshold,", indent=False)
    P.equation("adapt", math(v("T"), delim(cat(v("x"), t(","), v("y"))), t("="), nary("∑", cat(v("s"), t(","), v("t")), None,
                                                                                        cat(v("w"), delim(cat(v("s"), t(","), v("t"))), v("f"), delim(cat(v("x"), t("+"), v("s"), t(","), v("y"), t("+"), v("t"))))),
                             t("−"), v("C"), t(",")))
    P.para("with *w* a normalized 35×35 Gaussian window and *C* = 11; a pixel becomes white if *f* > *T*. For "
           "comparison, Otsu's method [@39] picks the global *t* that maximizes", indent=False)
    P.equation("otsu", math(sup(sub(v("σ"), v("B")), t("2")), delim(v("t")), t("="), sub(v("ω"), t("0")), delim(v("t")), sub(v("ω"), t("1")), delim(v("t")),
                            sup(delim(cat(sub(v("μ"), t("0")), delim(v("t")), t("−"), sub(v("μ"), t("1")), delim(v("t"))), "[", "]"), t("2")), t(",")))
    P.para("with class probabilities ω_{0}, ω_{1} and means μ_{0}, μ_{1}. Every mode ends with an 18 px white border, "
           "so OCR always receives a gray or binary page with a margin.", indent=False)


    P.subheading("E", "Transform Coding and Motion")
    P.para("Stored previews are JPEG files, that is, 8×8 block DCT coding [@53] with")
    P.equation("dct", math(v("F"), delim(cat(v("u"), t(","), v("v"))), t("="), frac(cat(v("C"), delim(v("u")), v("C"), delim(v("v"))), t("4")),
                           nary("∑", cat(v("x"), t("="), t("0")), t("7"), nary("∑", cat(v("y"), t("="), t("0")), t("7"),
                                cat(v("f"), delim(cat(v("x"), t(","), v("y"))), sub(v("c"), cat(v("x"), v("u"))), sub(v("c"), cat(v("y"), v("v")))))), t(",")))
    P.para("where *c*_{xu} = cos[(2*x* + 1)*u*π/16], *C*(0) = 1/√2 and *C*(*u*) = 1 otherwise [@52]; quantization "
           "rounds *F*/*Q* for a quality-scaled table *Q*, and PSNR = 10 log_{10}(255²/MSE). Block matching estimates "
           "a block's displacement *d* = (*u*, *v*), |*u*|, |*v*| ≤ *p*, between a current frame *c* and a reference "
           "*r* by minimizing [@56,@57]", indent=False)
    P.equation("mad", math(v("D"), delim(v("d")), t("="), frac(t("1"), sup(v("N"), t("2"))),
                           nary("∑", cat(v("p"), t("∈"), v("B")), None,
                                delim(cat(v("c"), delim(v("p")), t("−"), v("r"), delim(cat(v("p"), t("+"), v("d")))), "|", "|")), t(",")))
    P.para("over the pixels *p* of the *N*×*N* block *B*. Exhaustive search (ES) tests all (2*p* + 1)² = 225 offsets "
           "for *p* = 7; three-step search (TSS) [@57] tests 9 points at step 4 and 8 more at steps 2 and 1, 25 in "
           "all; 2-D logarithmic search [@56] moves a five-point cross and halves the step when the center wins. Both "
           "fast searches assume the cost falls monotonically toward the minimum. The app captures single stills, "
           "so motion estimation appears only in Pipeline Lab and the benchmark.", indent=False)

    # ---------------- V. Pipeline Lab
    P.heading("V", "Pipeline Lab")
    P.para("Pipeline Lab runs eleven stages on the current scan or a bundled synthetic certificate and shows the "
           "intermediate images, histograms and measurements (Fig. {fig:lab}): (1) spatial and "
           "gray-level resolution and PNG/JPEG sizes; (2) negative, contrast stretch, log, gamma "
           "and intensity slicing; (3) histograms, global equalization and CLAHE; (4) mean, Gaussian and median "
           "filtering of salt-and-pepper noise with PSNR, sharpening, the DFT spectrum and a Gaussian low-pass; "
           "(5) 8×8 DCT coding with a quality slider; (6) Sobel, Prewitt, Laplacian, LoG and Canny edges; (7) "
           "dilation, erosion, opening, closing and boundary extraction; (8) fixed, Otsu and adaptive thresholds with "
           "4- and 8-connected components; (9) RGB, HSV and YUV channels with hue slicing of ink and stamp; "
           "(10) page geometry with Euclidean, city-block and chessboard distances; and (11) block matching.")
    P.para("Point operations, histograms, Otsu, labelling, binary morphology, the DCT and the three searches are "
           "unit-tested Kotlin written from first principles; the other stages call OpenCV and the production "
           "*DocumentScanner* functions, so the Lab shows exactly what the scanner does. In stage 11, frame *t* is synthesized from frame *t* − 1 by a 4 px / −3 px shift, "
           "a 0.4° rotation and noise; ES, TSS and 2-D log search run on 16×16 blocks with *p* = 7, blank blocks are "
           "ignored, and a dominant vector longer than 1.5 px is reported as camera motion. This steady-capture "
           "decision is a Lab demonstration, not part of the capture screen.")
    P.figure("lab", FIG / "fig_lab_screens.png", COL_W,
             "Pipeline Lab on the synthetic sample: (a) corners, closed Canny edges, warp and enhancement; "
             "(b) input, equalized and CLAHE histograms; (c) points tested by ES, TSS and 2-D log search for one "
             "block, with the cost of each method.")

    # ---------------- VI. Recognition, layout and fields
    P.heading("VI", "Recognition, Layout and Field Extraction")
    P.para("ML Kit Text Recognition v2 reads printed text with its bundled Latin model [@35]. A formatter orders "
           "blocks and lines top to bottom and left to right and indents lines by zero to six spaces. In handwriting mode, each ML Kit line box is "
           "cropped from the enhanced page with margins of 4% (horizontal) and 22% (vertical), resized to a height of "
           "48 px with a width that is a multiple of 8 (16–2048 px), scaled to [−1, 1] and read by the 7.96 MB "
           "PP-OCRv5 English recognizer [@6] in ONNX Runtime [@33] with greedy CTC decoding [@10]; below a "
           "confidence of 0.20 the ML Kit line is kept. This path returns text only: correction and field "
           "extraction are skipped.")
    P.para("In printed and automatic modes, a conservative corrector maps digit–letter confusions and replaces a "
           "word only when a dictionary word is within Levenshtein distance 1 (2 for seven or more letters) [@81] and "
           "clearly closer than the runner-up; short, all-capital and digit-bearing words are protected. Field "
           "extraction is a cascade. Label rules (\"Issued at *place* on *date*\", labelled dates, port lines, header "
           "organizations) run first. Missing fields go to a MiniLM question-answering model fine-tuned on "
           "certificate text and exported to int8 ONNX (32.3 MiB, 384-token windows, stride 96); if it cannot run, the "
           "100.46 MB MobileBERT model [@28] answers. A model is queried only when matching evidence (for example "
           "expiry wording) is present, answers must be 2–160 characters, and dates must match a date pattern. "
           "MobileBERT also answers free questions over saved documents using 1,200-character chunks with "
           "180-character overlap.")
    P.para("Fig. {fig:screens} shows the main screens on an Android emulator (Pixel profile) with a synthetic "
           "certificate. The preview reports the quality verdict and whether page borders were found, and offers "
           "the Auto, Clean and B&W enhancement modes. An earlier build was also installed and run on an OPPO "
           "CPH2665 phone with Android 15.")
    P.figure("screens", FIG / "fig_app_screens.png", COL_W,
             "App screens with a synthetic certificate: (a) manual crop of a tilted photo; (b) aligned preview with "
             "quality verdict and enhancement modes; (c) extracted certificate details.")

    P.para("The adaptation corpus had 108 PDFs and 351 pages (295 unique), split by document into 206/56/33 "
           "pages. A 4.37M-parameter tiny BERT reached only 23.5% test exact match, so MiniLM, starting from a "
           "SQuAD 2.0 checkpoint [@29,@31], was fine-tuned on 169 weak rows from an expanded set of archives "
           "(Table {tab:qa}). Its scores are text-only pilot results on a six-document, 26-field gold set checked by "
           "hand; they exclude capture and OCR errors.")
    P.table("qa", "Certificate-field QA models (text input, exact match)",
            ["Model", "Weak val.", "Gold (26)", "Size"],
            [["Tiny BERT (4.37M)", "35.3%¹", "4 (15.4%)", "17.5 MB"],
             ["MiniLM, fp32", "35/49 (71.4%)", "15 (57.7%)", "126.9 MiB"],
             ["MiniLM, int8 (app)", "33/49 (67.3%)", "15 (57.7%)", "32.3 MiB"]],
            [1.2, 0.85, 0.75, 0.7],
            notes="¹Own validation split. Gold by field (MiniLM fp32): name 2/6, issue date 3/6, port 4/4, authority "
                  "2/4, expiry 4/6. The int8 model changed three individual gold predictions.")

    # ---------------- VII. Experimental setup
    P.heading("VII", "Experimental Setup")
    P.para("Real certificates are private, so the benchmark renders synthetic ones whose text, fields and corners "
           "are known; no real certificate appears in this paper. Forty A4 pages (1240×1754 px, seeds 1000–1039) use "
           "one of 12 fonts, a title, a body paragraph, seven labelled fields in three date formats, a decorative "
           "double border, a colored seal and a signature. Each page is warped into an 1800×2400 frame on a wood, "
           "laminate or fabric background with clutter, drop shadow and vignetting, with 3% corner jitter, noise "
           "σ = 2 and JPEG quality 92. One degradation is then applied at three severities: tilt (jitter "
           "0.06/0.10/0.14), cast shadow (0.65/0.45/0.30), low light (gain 0.40/0.25/0.15), overexposure (gain "
           "1.45/1.80/2.20), Gaussian blur (σ = 1.2/2.0/3.0 px), motion blur (7/13/19 px), noise (σ = 10/20/32) or "
           "JPEG (quality 25/12/6). With a clean condition this gives 25 conditions × 8 captures, *n* = 200.")
    P.para("A line-by-line Python port of *DocumentScanner.kt* yields four variants per capture: raw photo, "
           "rectified, rectified + Auto, and rectified + B&W. OCR is PaddleOCR 3.7.0 with PP-OCRv6 small detection "
           "and recognition [@8] in ONNX Runtime 1.30.0 on the CPU, with orientation and unwarping disabled. "
           "After lower-casing and collapsing whitespace,")
    P.equation("cer", math(t("CER"), t("="), frac(cat(v("S"), t("+"), v("D"), t("+"), v("I")), v("N")), t(",")))
    P.para("the Levenshtein distance over the reference length; WER is the same over words, and both count "
           "reading-order errors, whereas word recall ignores order. Field exact match takes the first line containing "
           "a field label and compares the rest of it with the true value (5 fields, 1,000 slots). Page IoU uses the "
           "convex intersection of detected and true quads. The overall paired interval is the benchmark's bootstrap; "
           "per-degradation intervals are percentile bootstraps over captures (10,000 resamples, seed 152) computed "
           "from the per-capture results. Timings are desktop medians on an AMD Ryzen 9 8945H.", indent=False)

    # ---------------- VIII. Results
    P.heading("VIII", "Results")
    P.subheading("A", "Rectification and Enhancement")
    P.table("overall", "OCR on the same 200 captures (PP-OCRv6 small)",
            ["Variant", "Mean CER", "Median CER", "WER", "Word recall", "Field EM"],
            [["Raw photo", "9.6%", "0.49%", "14.1%", "92.9%", "69.1%"],
             ["Rectified", "4.2%", "0.10%", "8.7%", "92.4%", "88.4%"],
             ["+ Auto (gray)", "4.2%", "0.10%", "7.4%", "93.6%", "88.4%"],
             ["+ B&W adaptive", "7.6%", "0.58%", "14.6%", "86.7%", "74.0%"]],
            [0.9, 0.52, 0.55, 0.42, 0.55, 0.52])
    P.para("Rectification more than halves the mean CER and raises field exact match from 691 to 884 of 1,000 slots "
           "(Table {tab:overall}). Against the raw photo, the Auto page lowers CER by 5.4 points (95% CI 3.9–7.1), "
           "better on 112 captures, worse on 40 and tied on 48. The gain is geometric: rectified and enhanced pages "
           "have the same CER and field score, and word recall barely moves (92.9% to 93.6%). Every field improves "
           "with rectification (of 200): certificate name 136→177, issue date 145→180, port 152→180, authority "
           "117→170, expiry 141→177. B&W is worse than gray on average.")
    P.figure("cer", FIG / "fig_cer_field_by_degradation.png", COL_W,
             "Mean CER (top) and five-field exact match (bottom) by degradation. Error bars: 95% percentile-bootstrap "
             "intervals over captures.")
    rows = []
    names = {"clean": "Clean", "tilt": "Tilt", "shadow": "Shadow", "low_light": "Low light", "overexposure": "Overexp.",
             "gaussian_blur": "Gauss. blur", "motion_blur": "Motion blur", "noise": "Noise", "jpeg": "JPEG"}
    det = {"clean": "100/100", "tilt": "100/100", "shadow": "100/100", "low_light": "33/29",
           "overexposure": "100/46", "gaussian_blur": "100/96", "motion_blur": "100/46",
           "noise": "100/100", "jpeg": "100/100"}
    for k, nm in names.items():
        e = S[k]
        dlt = e["paired_cer_raw_minus_gray"]
        rows.append([nm, det[k], pct(e["cer_raw"][0]), pct(e["cer_rectified"][0]), pct(e["cer_enhanced_gray"][0]),
                     f"{pct(dlt[0])} ({pct(dlt[1])}, {pct(dlt[2])})",
                     pct(e["field_em_raw"][0]), pct(e["field_em_enhanced_gray"][0])])
    P.subheading("B", "Effect by Degradation")
    P.table("deg", "Results by degradation (%), n = 24 per row (clean 8)",
            ["Degradation", "Quad / IoU≥0.9", "CER raw", "CER rect.", "CER gray", "ΔCER raw−gray (95% CI)", "EM raw", "EM gray"],
            rows, [0.62, 0.48, 0.32, 0.34, 0.34, 0.78, 0.3, 0.32], size=7,
            notes="Quad: page quadrilateral found. ΔCER: paired reduction from raw photo to Auto gray. EM: five-field "
                  "exact match.")
    P.para("Fig. {fig:cer} and Table {tab:deg} break the result down. Tilt dominates: raw CER of 20.3% falls to 0.4% "
           "after rectification although raw word recall is already 99.4%, so the words are read but the reading "
           "order is scrambled. Enhancement matters where detection fails or text is soft. Under low light a quad is "
           "found on only 33% of captures; at the strongest level, where raw and rectified are the same image, CER "
           "falls from 11.1% to 7.2% with Auto and to 3.8% with B&W. Under Gaussian blur, word recall rises from "
           "90.8% to 96.3%. Under overexposure enhancement hurts (rectified 1.1%, Auto 2.7%), plausibly because CLAHE and "
           "sharpening amplify the clipped background. Motion blur is the hardest case (24.4% after rectification), "
           "and binarization worsens it to 41.5% by breaking smeared strokes.")

    P.subheading("C", "Page Detection and Quality Checks")
    P.para("A quadrilateral was found on 184 of 200 captures and matched the true page with IoU ≥ 0.9 on 156; when "
           "found, mean IoU was 0.976 and median corner error 1.3 px in the 1800×2400 frame. Misses concentrate in "
           "low light; IoU drops under overexposure and motion blur.")
    P.table("quality", "Quality verdict vs applied degradation (n = 200)",
            ["Expected \\ app label", "Good", "Dark", "Bright", "Soft"],
            [["Good (n = 104)", "104", "0", "0", "0"], ["Too dark (n = 24)", "8", "16", "0", "0"],
             ["Too bright (n = 24)", "5", "0", "19", "0"], ["Soft image (n = 48)", "37", "0", "0", "11"]],
            [1.3, 0.5, 0.5, 0.55, 0.5])
    P.para("The quality checks raised no false alarm and gave the expected label on 150 of 200 captures (Table "
           "{tab:quality}). The dark flag fired on every medium and strong low-light capture; the eight misses were "
           "mild ones with mean gray 65. The blur flag caught all eight strong Gaussian-blur captures but only 1 of 24 "
           "motion-blurred ones: motion blur keeps strong second derivatives across the stroke direction, so the "
           "whole-frame Laplacian variance, background texture included, stays above 65.")

    P.subheading("D", "Thresholding Under Shadow")
    P.figure("thr", FIG / "fig_threshold_comparison.png", COL_W,
             "Global, Otsu and adaptive thresholds on a strongly shadowed page (top) and mean CER over the 24 shadow "
             "captures (bottom).")
    P.table("thr", "Binarization of 24 shadowed pages (true corners)",
            ["Method", "CER gray", "Ink F1", "CER app path", "Ink F1"],
            [["Global *T* = 127", "34.2%", "0.41", "33.6%", "0.30"],
             ["Otsu", "39.2%", "0.25", "40.2%", "0.22"],
             ["Adaptive 35/11", "0.05%", "0.91", "0.40%", "0.87"],
             ["No threshold", "0.03%", "–", "–", "–"]],
            [1.1, 0.6, 0.5, 0.75, 0.5],
            notes="App path: after CLAHE and 3×3 median. Ink F1: pixel F1 against the rendered ink mask.")
    P.para("Rectifying the 24 shadow captures with their true corners isolates binarization (Fig. {fig:thr}, Table "
           "{tab:thr}). Only the adaptive threshold survives, with CER 0.05%, close to the 0.03% of the gray page. "
           f"Otsu fails even under the mildest shadow ({pct(thr['otsu_gray']['by_severity_cer']['1'])}% CER), and the "
           f"fixed threshold works at mild severity ({pct(thr['global127_gray']['by_severity_cer']['1'], 2)}%) but not "
           f"at medium or strong ({pct(thr['global127_gray']['by_severity_cer']['2'])}% and "
           f"{pct(thr['global127_gray']['by_severity_cer']['3'])}%). CLAHE and the median do not rescue either global "
           "method.")

    P.subheading("E", "DCT Compression")
    P.table("dct", "8×8 DCT coding of six clean pages",
            ["Quality factor", "90", "75", "50", "30", "20", "10", "5", "2"],
            [["PSNR (dB)", "46.0", "39.0", "34.9", "32.3", "30.3", "27.7", "25.4", "23.6"],
             ["Nonzero coeff. (%)", "12.7", "10.1", "8.1", "6.8", "6.0", "4.7", "3.8", "3.1"],
             ["JPEG bit/pixel", "0.99", "0.72", "0.56", "0.45", "0.38", "0.29", "0.23", "0.19"],
             ["OCR CER (%)", "0.11", "0.41", "0.33", "0.28", "0.34", "0.35", "0.48", "1.21"]],
            [1.06] + [0.305] * 8,
            notes="Uncompressed CER 0.02%. Keeping only the top-left *k*×*k* coefficients gives CER 0.12%, 0.30%, "
                  "0.75%, 4.1% and 98.9% for *k* = 6, 4, 3, 2, 1 (PSNR 30.3 to 16.6 dB).")
    P.para("Text is robust to transform coding (Table {tab:dct}). At quality 10 only 4.7% of coefficients are "
           "nonzero and a real JPEG needs 0.29 bit/pixel, yet CER stays at 0.35% while PSNR falls to 27.7 dB. CER is "
           "not monotonic at high quality (0.41% at 75, 0.33% at 50), so differences below about half a point are "
           "noise. Zonal coding fails only below 3×3 coefficients, when strokes merge into block averages. The "
           "app's quality-88 previews therefore cost little in legibility.")

    P.subheading("F", "Block-Matching Motion Estimation")
    P.table("bm", "Block matching on 10 text frame pairs (9,990 blocks)",
            ["Frames, method", "Points/block", "Exact MV", "MV error (px)", "Mean MAD"],
            [["Sharp, ES", "225", "99.9%", "0.00", "1.60"],
             ["Sharp, TSS", "25", "29.6%", "5.05", "6.59"],
             ["Sharp, 2-D log", "20.7 (≤ 31)", "32.8%", "5.10", "6.82"],
             ["Low-pass, ES", "225", "92.9%", "0.10", "0.22"],
             ["Low-pass, TSS", "25", "49.1%", "3.22", "0.82"],
             ["Low-pass, 2-D log", "21.5 (≤ 31)", "58.7%", "2.96", "0.72"]],
            [1.05, 0.72, 0.58, 0.68, 0.55],
            notes=f"Exact MV and MV error over textured blocks (std. dev. > 10): "
                  f"{bmj['sharp_frames']['textured_blocks']:,} sharp, {bmj['lowpass_sigma2_frames']['textured_blocks']:,} "
                  "low-pass. Low-pass: Gaussian σ = 2 on both frames.")
    P.para("Ten pairs of 480×640 frames were made by shifting a synthetic capture by a random integer vector in "
           "[−7, 7]² and adding noise (σ = 2), so every true vector is known. ES recovers it on 99.9% of textured "
           "blocks (Table {tab:bm}). TSS and 2-D log search test 9–11 times fewer points but are exact on only 29.6% "
           "and 32.8% of blocks, with mean errors above 5 px. Fig. {fig:bm} shows why: on text the cost surface has "
           "many local minima, and the first coarse step commits to the wrong one. Low-pass filtering both frames "
           "raises the fast searches to 49.1% and 58.7%, at a small cost to ES (92.9%).")
    P.figure("bm", FIG / "fig_block_matching.png", COL_W,
             "Points visited by (a) TSS and (b) 2-D log search for one text block; both settle at the local "
             "minimum (2, 4), while the true vector is (−3, 4).")

    P.subheading("G", "Processing Time")
    P.para("Median desktop times per 1800×2400 capture (12 captures × 5 repeats) were 1.3 ms (resize), 14.6 ms "
           "(quality check), 6.3 ms (detection), 8.8 ms (warp) and 6.2 ms or 19.1 ms (Auto or B&W enhancement), "
           "against 1063 ms for one PP-OCRv6 call: the classical Auto path costs about 37 ms, roughly 3.5% of OCR. "
           "Phone latency was not measured.")
    P.heading("IX", "Discussion")
    P.para("*Why rectification dominates.* Learned line detectors tolerate skew, so on a tilted page "
           "nearly every word is still read; what breaks is the line structure: slanted lines interleave, and a label "
           "no longer shares a line with its value. The homography restores horizontal lines and with them the "
           "label–value structure that label-based extraction needs, which is why CER and field match move together "
           "while word recall stays flat.")
    P.para("*When enhancement helps and hurts.* A learned recognizer already normalizes contrast, so CLAHE and "
           "sharpening add little on well-exposed pages. They help where the signal is weak (dark or Gaussian-blurred "
           "captures, especially when no page is found) and hurt on clipped input, where CLAHE stretches residual "
           "texture and the unsharp mask overshoots. The measured brightness and sharpness could pick the mode.")
    P.para("*Why Otsu fails under shadow.* A cast shadow splits the paper into two background modes, so the "
           "variance-maximizing threshold separates lit from shadowed paper instead of ink from paper; a local "
           "threshold follows the illumination.")
    P.para("*Why fast block matching fails on text.* Text is a quasi-periodic texture: a shift of one character "
           "or line pitch gives a near-equal match, so the cost surface is not unimodal and a coarse step lands in "
           "the wrong basin. Smoothing widens the true basin. For a steady-capture trigger, exhaustive search on a "
           "small, smoothed preview frame is the safer choice.")
    P.para("*Implications for the app.* Rectification should stay mandatory and Auto gray the default; B&W suits "
           "uneven light but not blur or noise. The quality check should run on the detected page and needs a "
           "directional blur measure. In a companion synthetic experiment, a learned depth-3 decision tree on the "
           "same two features reached 92.7% accuracy against 78.1% for the hand-set thresholds, which supports "
           "calibrating them on real phone captures.")
    P.heading("X", "Limitations and Future Work")
    P.para("The benchmark is synthetic: clean one-column pages with known fonts and a simple camera model. Real "
           "captures add lens distortion, curl, glare, handwriting, tables and stamps over text, so the absolute "
           "numbers are optimistic; the paired differences are the result. The benchmark OCR is PP-OCRv6 on a "
           "desktop while the app uses ML Kit, so the study must be repeated on the phone with latency and memory. Samples are small (24 captures per degradation, 26 gold fields), and the "
           "benchmark's field rule is simpler than the app's cascade.")
    P.para("Contour detection fails on borderless and low-contrast pages; corner ordering by *x* + *y* and "
           "*y* − *x* can collide for a page rotated near 45°; the Canny and quality thresholds need per-device "
           "calibration. The companion evaluation also exposed two post-correction defects, a digit guard that kept the "
           "confusion map from firing and a dictionary without maritime words; both are now fixed and locked by "
           "regression tests. Future work will build a "
           "vessel-disjoint set of real photographed certificates with verified text, fields and corners; try learned "
           "detection and dewarping [@11,@21,@22] where contours fail; compare PP-OCRv6 with ML Kit on the "
           "phone; add an expiry-after-issue consistency check; and test layout-aware [@16] and OCR-free [@14] "
           "extractors within the mobile budget.")

    # ---------------- XI. Conclusion
    P.heading("XI", "Conclusion")
    P.para("NavCert Offline captures, rectifies, enhances, reads and parses a certificate entirely on a phone. On "
           "200 paired synthetic captures, its classical front end lowers mean CER from 9.6% to 4.2% and lifts "
           "five-field exact match from 69.1% to 88.4%, almost all of it from geometric rectification. Adaptive "
           "thresholding is the right binarizer under shadow, DCT coding at preview quality costs almost nothing in "
           "legibility, and fast block matching is unreliable on text unless frames are smoothed first. Pipeline Lab "
           "makes each operation visible on the user's own scan. The next step is the same paired measurement on "
           "real certificates photographed with the phone.")

    # ---------------- references
    P.heading("", "References")
    for i, n in enumerate(P.cite_order, 1):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.left_indent = Inches(0.28)
        pf.first_line_indent = Inches(-0.28)
        pf.space_after = Pt(0.5)
        pf.line_spacing = Pt(9)
        pf.tab_stops.add_tab_stop(Inches(0.28))
        r = p.add_run(f"[{i}]\t"); set_font(r, 8)
        for k, seg in enumerate(re.split(r"\*", P.refs[n])):
            if seg:
                r = p.add_run(seg); set_font(r, 8, italic=bool(k % 2))
    P.section(2)

    props = d.core_properties
    props.title = TITLE
    props.author = "Aadidev Raizada"
    props.subject = "Image and Video Processing research paper"
    props.keywords = "document image processing, rectification, CLAHE, adaptive thresholding, mobile OCR, block matching"
    return P


def main():
    pre = build().labels          # pass 1: collect figure/table/equation numbers
    P = build(pre)                # pass 2: resolve forward references
    P.doc.save(DOCX)
    print("references cited:", len(P.cite_order), "figures:", P.fig_n, "tables:", P.tab_n, "equations:", P.eq_n)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(DOCX)], check=True,
                   capture_output=True)
    print(PDF)


if __name__ == "__main__":
    main()
